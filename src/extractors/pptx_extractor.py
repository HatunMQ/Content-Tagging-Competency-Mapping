"""
Extract text chunks from a .pptx file — one chunk per slide.

Native text (titles, text boxes, tables) is pulled first. Optional OCR
(pytesseract) only kicks in on a slide that has little/no native text,
which is exactly the "text baked into an image/diagram" case.

Before OCR runs, a picture that looks like a full desktop/browser
screen capture (not a clean diagram — e.g. someone screen-recorded a
YouTube video instead of cropping just the relevant frame) gets its
browser chrome / sidebar / taskbar cropped out first. Without this,
OCR reads browser tab titles, ads, and the taskbar clock as if they
were slide content — verified directly on Decision Trees Revised.pptx
slide 38 (the "Free Email Addresses / WhatsApp / mantasleep.com" junk
came from OCR reading the whole desktop screenshot, not the diagram).

lang="eng" only, not "eng+ara": this content is English-only, so
Arabic is not offered to Tesseract as a candidate alphabet. Giving OCR
"eng+ara" on English-only slides invites it to force-fit ambiguous
shapes (icons, logos, low-res UI elements) into Arabic glyphs, which
is where the stray Arabic-looking characters in eval_sample.csv came
from.

If pytesseract or its system Tesseract binary aren't installed, OCR is
silently skipped — the script still works, it just relies on native text.
"""
import io
import re
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

try:
    import pytesseract
    from PIL import Image
    OCR_AVAILABLE = True
except ImportError:
    OCR_AVAILABLE = False

# Crop margins for pictures that look like full-screen captures, tuned
# against the real screenshot found in this project's data (1920x1080,
# Chrome browser + Windows taskbar). Adjust if your screen-recordings
# use a different browser/OS chrome layout.
CROP_TOP_PCT = 0.10      # browser tab bar + address bar
CROP_RIGHT_PCT = 0.32    # e.g. YouTube "recommended videos" sidebar
CROP_BOTTOM_PCT = 0.05   # OS taskbar

BIDI_MARKS = re.compile(r"[‎‏‪-‮⁦-⁩]")
ARABIC_RUN = re.compile(r"[؀-ۿ]+")
WHITESPACE = re.compile(r"[ \t]+")


def _clean_ocr_text(text: str) -> str:
    text = BIDI_MARKS.sub("", text)
    text = ARABIC_RUN.sub("", text)  # safety net — this content is English-only
    text = WHITESPACE.sub(" ", text)
    return text.strip()


def _looks_like_screen_capture(width_px: int, height_px: int) -> bool:
    """A full desktop/browser screen recording is usually a standard
    screen resolution with a ~16:9 aspect ratio — unlike a purposely
    cropped diagram or chart image."""
    if height_px == 0:
        return False
    ratio = width_px / height_px
    common_res = {(1920, 1080), (1366, 768), (2560, 1440), (1280, 720)}
    return (width_px, height_px) in common_res or (1.7 <= ratio <= 1.8 and width_px >= 1200)


def _extract_slide_text(slide) -> str:
    parts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
        if shape.has_table:
            for row in shape.table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip(" |"):
                    parts.append(row_text)
    return "\n".join(parts)


def _ocr_slide_images(slide) -> list[str]:
    if not OCR_AVAILABLE:
        return []
    texts = []
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                image_bytes = shape.image.blob
                img = Image.open(io.BytesIO(image_bytes))
                w, h = img.size
                if _looks_like_screen_capture(w, h):
                    top = int(h * CROP_TOP_PCT)
                    right = int(w * (1 - CROP_RIGHT_PCT))
                    bottom = int(h * (1 - CROP_BOTTOM_PCT))
                    img = img.crop((0, top, right, bottom))
                ocr_text = pytesseract.image_to_string(img, lang="eng").strip()
                ocr_text = _clean_ocr_text(ocr_text)
                if ocr_text:
                    texts.append(ocr_text)
            except Exception:
                continue
    return texts


def extract_pptx(path: Path, use_ocr: bool = True) -> list[dict]:
    prs = Presentation(str(path))
    chunks = []

    for i, slide in enumerate(prs.slides, start=1):
        text = _extract_slide_text(slide)

        # Only pay the (slow) OCR cost when native text is thin — most slides
        # already have enough text and don't need it.
        if use_ocr and len(text) < 40:
            ocr_texts = _ocr_slide_images(slide)
            if ocr_texts:
                text = (text + "\n" + "\n".join(ocr_texts)).strip()

        if not text.strip():
            continue  # skip pure-decoration slides (section dividers, logos)

        title = (
            slide.shapes.title.text.strip()
            if slide.shapes.title and slide.shapes.title.text
            else f"Slide {i}"
        )

        chunks.append({
            "file_name": path.name,
            "chunk_id": f"slide_{i:02d}",
            "content_type": "Visual lecture",
            "text": text,
            "metadata": {"source_location": f"slide {i}", "section_title": title},
        })

    return chunks


if __name__ == "__main__":
    import sys
    import json

    p = Path(sys.argv[1])
    for c in extract_pptx(p):
        print(json.dumps(c, ensure_ascii=False))